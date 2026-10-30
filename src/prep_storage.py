"""
prep_storage.py
Owner: Rongjia

PURPOSE
    Turn the storage estimates (EIA working-gas capacity and March 2025
    inventories, National Fuel FSS fees, UGI/Crestwood injection and withdrawal
    rates) into one clean CSV the model can read. This script reads your final
    estimates from data/raw/storage/ and reshapes them.

INPUTS
    data/raw/storage/     Source files, your estimation notebook/Excel, and the
                          final storage parameter table, committed as-is.

OUTPUT
    data/processed/storage.csv    2 rows, one per leased field, these columns:

    field            str    H (hub, PA) | M (market area, NY)
    market           str    HUB for H, NY for M
    capacity_dth     float  leased working-gas capacity, Dth        (Ibar_f)
    init_inv_dth     float  inventory at start of April 2025, Dth   (I_f,0)
    max_inj_dth_day  float  maximum injection rate, Dth/day         (r_in_f)
    max_wd_dth_day   float  maximum withdrawal rate, Dth/day        (r_out_f)
    inj_fee          float  injection charge, $/Dth                 (c_in_f)
    wd_fee           float  withdrawal charge, $/Dth                (c_out_f)

    Must hold: 0 <= init_inv_dth <= capacity_dth.

ALSO OWED
    data/processed/ranges/storage_ranges.csv with columns
    parameter, base, low, high, rationale
    (at minimum the leased shares behind capacity_dth, and the rate limits),
    for Milestone 5. Report section is shared with prep_demand.py.

RULES
    - Run from the repo root:  python src/prep_storage.py
    - Relative paths only. Units: Dth and $/Dth.
    - Re-running must reproduce storage.csv exactly. Never edit the CSV by hand.

DUE: Thu Oct 29 (pull request to main; Rob reviews)
"""



from pathlib import Path
import csv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

INPUT_FILE = (
    ROOT / "data" / "raw" / "storage"
    / "The daily withdraw and injection limits and costs.xlsx"
)

OUTPUT_FILE = ROOT / "data" / "processed" / "storage.csv"
RANGES_FILE = (
    ROOT / "data" / "processed" / "ranges"
    / "storage_ranges.csv"
)

FIELDS = [
    "field", "market", "capacity_dth", "init_inv_dth",
    "max_inj_dth_day", "max_wd_dth_day", "inj_fee", "wd_fee"
]

LEASE_SHARES = {"H": 0.005, "M": 0.01}

# STORAGE ASSUMPTIONS
# Source: EIA March 2025 storage capacity and inventory.
# Leased capacity shares: PA 0.5%, NY 1.0%.
# Initial inventory uses historical state fill ratios.
#
# Injection benchmarks: UGI Storage (PA), Thomas Corners (NY).
# Withdrawal benchmarks: EIA-191 field-level data.
# Operating fees: National Fuel 2025 FSS tariff, $0.0381/Dth.
#
# Operating limits are scaled to assumed leased capacity.
# Rates are constant; inventory-dependent limits are excluded.
# See writeups/demand_storage_assumptions.md for details.
# This script reads final benchmark-based estimates from Excel.
# Original EIA and facility data are retained for verification.
# The script does not recalculate facility-level benchmarks.
def main() -> None:
    wb = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    try:
        ws = wb["Data"]
        records = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            name = str(row[0]).strip()
            if name.startswith("H"):
                field, market = "H", "HUB"
            elif name.startswith("M"):
                field, market = "M", "NY"
            else:
                raise ValueError(f"Unknown storage facility: {name}")

            values = [float(v) for v in row[1:7]]
            capacity, inventory, inj, wd, inj_fee, wd_fee = values

            assert 0 <= inventory <= capacity
            assert capacity > 0 and inj > 0 and wd > 0
            assert inj_fee >= 0 and wd_fee >= 0

            records.append(dict(zip(FIELDS, [
                field, market, *values
            ])))
    finally:
        wb.close()

    assert len(records) == 2
    assert {r["field"] for r in records} == {"H", "M"}

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_FILE.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(records)

    # Preliminary sensitivity scenarios (not confidence intervals).
    ranges = []

    for r in records:
        field = r["field"]
        parameters = {
            f"lease_share_{field}": LEASE_SHARES[field],
            f"max_inj_dth_day_{field}": r["max_inj_dth_day"],
            f"max_wd_dth_day_{field}": r["max_wd_dth_day"],
        }

        for parameter, base in parameters.items():
            ranges.append({
                "parameter": parameter,
                "base": base,
                "low": round(base * 0.8, 8),
                "high": round(base * 1.2, 8),
                "rationale": "Assumed +/-20% sensitivity scenario"
            })

    RANGES_FILE.parent.mkdir(parents=True, exist_ok=True)
    with RANGES_FILE.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f, fieldnames=["parameter", "base", "low", "high", "rationale"]
        )
        writer.writeheader()
        writer.writerows(ranges)

    print(f"Created {OUTPUT_FILE}: {len(records)} rows")
    print(f"Created {RANGES_FILE}: {len(ranges)} rows")




if __name__ == "__main__":
    main()
